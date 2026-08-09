from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from app.services.excel_service import (
    _to_numeric,
    find_column,
    list_excel_files,
    read_excel_source,
)

GUIDE_DIR_CANDIDATES = ("guide-bo-report", "guid-bo-report")
GUIDE_FILENAME = "Guide BO Report.xlsx"
GUIDE_SHEET = "Guide"
TEMPLATE_SHEET = "Data Template"
ESTIMASI_SHEET = "Estimasi & Remark"
TEMPLATE_NO_PO_SHEET = "Template NO PO"
GUIDE_KEEP_SHEETS = {
    TEMPLATE_SHEET,
    GUIDE_SHEET,
    ESTIMASI_SHEET,
    TEMPLATE_NO_PO_SHEET,
}

# Template NO PO header → possible Data Template column names (first match wins)
NO_PO_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "Plant": ("Plant",),
    "Customer Name": ("Customer Name",),
    "Customer PO": ("Customer PO",),
    "IREQ Item": ("IREQ Item",),
    "SO Date": ("SO Date",),
    "Sales document": ("Sales document",),
    "Sales Document Item": ("Sales Document Item",),
    "Class": ("Class",),
    "Created By": ("Created by", "Created By"),
    "Material No": ("Material No",),
    "Material Description": ("Material Description",),
    "Order Quantity": ("Order Quantity",),
    "SNSKI": ("SNSKI",),
    "Gross Weight": ("Gross Weight", "Weight"),
    "On-hand Stock": ("On-hand Stock", "SOH"),
    "On-Order Qty": ("On-Order Qty", "On-Order"),
    "Pre-Stock Qty": ("Pre-Stock Qty", "Pre-Stock"),
    "On Hand Reserve Stock": ("On Hand Reserve Stock", "Reserve"),
    "1G38": ("1G38",),
    "1S67": ("1S67",),
    "1S66": ("1S66",),
    "1S76": ("1S76",),
    "1S81": ("1S81",),
    "Remark": ("Remark", "Remaks"),
    "Order Method": ("Order Method",),
    "Purchasing Document": ("Purchasing Document",),
    "Purchase Requisition": ("Purchase Requisition",),
    "PR Item": ("PR Item",),
    "Need By Date": ("Need By Date",),
    "Deletion indicator": ("Deletion indicator",),
}
OUTPUT_BO_REPORT = "bo_report.xlsx"
JOB_ID_BO_GUIDE = "bo-guide"
EXCLUDE_MATERIAL_NOS = {"DELIVERY_CHARGE:ZZ"}

# Stock Info plant 1G38 → Data Template.
# Guide uses SAP titles (Gross Weight, On-hand Stock, ...); Data Template may use
# short aliases (Weight, SOH, ...). Fill whichever header exists in the template.
STOCK_1G38_FIELD_SPECS: list[tuple[tuple[str, ...], tuple[str, ...]]] = [
    (("SNSKI",), ("SNSKI",)),
    (("Weight", "Gross Weight"), ("Gross Weight", "Weight")),
    (("SOH", "On-hand Stock"), ("On-hand Stock", "SOH")),
    (("On-Order", "On-Order Qty"), ("On-Order Qty", "On-Order")),
    (("Pre-Stock", "Pre-Stock Qty"), ("Pre-Stock Qty", "Pre-Stock")),
    (("Reserve", "On Hand Reserve Stock"), ("On Hand Reserve Stock", "Reserve")),
    (("Hazardous Indicator",), ("Hazardous Indicator",)),
    (("Replacement Indicator",), ("Replacement Indicator",)),
    (("Deletion Indicator",), ("Deletion Indicator",)),
    (("Returnable Indicator",), ("Returnable Indicator",)),
    (("Package Qty",), ("Package Qty",)),
    (("Commodity Code",), ("Commodity Code",)),
]

# Guide: 1G38 and hub plants take On-hand Stock (ATP)
STOCK_QTY_CANDIDATES = (
    "On-hand Stock (ATP)",
    "On-hand Stock",
    "Total Availability (TA)",
)

DATE_OUTPUT_COLUMNS = {
    "SO Date",
    "Released Date",
    "Need By Date",
    "PO Created Date",
    "Ship Out",
    "Ship In",
    "Invoice Date",
    "Old ESD",
    "Est Ship Date",
    "Shipment Number Date",
    "Act Dept Dt -Inv CAT",
    "Shp By Dt",
    "Source Date",
    "Act GI Date of OD",
    "Gate Pass Print Date",
    "Estimasi",
    "Estimasi before",
}

FULL_OD_GROUP_COLUMNS = [
    "Sales document",
    "Sales Document Item",
    "Material No",
    "Order Quantity",
]

DEDUPE_KEY_COLUMNS = [
    "Sales document",
    "Sales Document Item",
    "Material No",
    "Order Quantity",
    "OD of Sales Order Item",
    "OD Item of Sales Order Item",
    "OD Quantity",
]

NUMERIC_OUTPUT_COLUMNS = {
    "Total Price",
    "Order Quantity",
    "PO Quantity",
    "SNG",
    "Mell",
    "QNS",
    "SAG",
    "ETA D",
    "1S67",
    "1S66",
    "1S76",
    "1S81",
    "OD Quantity",
    "1G38",
    "Weight",
    "Gross Weight",
    "SOH",
    "On-hand Stock",
    "On-Order",
    "On-Order Qty",
    "Pre-Stock",
    "Pre-Stock Qty",
    "Reserve",
    "On Hand Reserve Stock",
    "Package Qty",
}


def _norm(series: pd.Series) -> pd.Series:
    return (
        series.astype(str)
        .str.strip()
        .replace({"": pd.NA, "nan": pd.NA, "None": pd.NA, "NaT": pd.NA})
    )


def _material_base(series: pd.Series) -> pd.Series:
    """Strip ':suffix' from material numbers (3574260:AA -> 3574260)."""
    return _norm(series).astype(str).str.split(":", n=1).str[0]


def _qty_key(series: pd.Series) -> pd.Series:
    num = _to_numeric(series)
    as_int = num.round(0)
    text = as_int.apply(lambda v: str(int(v)) if pd.notna(v) else "")
    return text


def _safe_col(df: pd.DataFrame, candidates: list[str]) -> str | None:
    try:
        return find_column(df, candidates)
    except KeyError:
        return None


def _template_targets(
    template_headers: list[str],
    aliases: tuple[str, ...],
) -> list[str]:
    """Return alias names that exist on the Data Template (or first alias as fallback)."""
    header_set = set(template_headers)
    found = [name for name in aliases if name in header_set]
    return found if found else [aliases[0]]


def _map_by_material_base(
    base_mat_base: pd.Series,
    keys: pd.Series,
    values: pd.Series,
) -> pd.Series:
    mapping: dict[str, str] = {}
    for k, v in zip(keys.tolist(), values.tolist(), strict=False):
        if k and k not in mapping:
            mapping[k] = v
    return base_mat_base.map(
        lambda k, m=mapping: m.get(str(k).strip(), "") if pd.notna(k) else ""
    )


def _load_folder_frames(folder: Path) -> pd.DataFrame:
    paths = list_excel_files(folder)
    if not paths:
        return pd.DataFrame()
    frames = [read_excel_source(path) for path in paths]
    return pd.concat(frames, ignore_index=True)


def _lookup_map(
    df: pd.DataFrame,
    key_series: pd.Series,
    value_col: str,
) -> dict[str, str]:
    if df.empty or value_col not in df.columns:
        return {}
    keys = _norm(key_series).fillna("")
    values = _norm(df[value_col]).fillna("")
    out: dict[str, str] = {}
    for key, value in zip(keys.tolist(), values.tolist(), strict=False):
        if not key or key in out:
            continue
        out[key] = value
    return out


def _map_series(keys: pd.Series, mapping: dict[str, str]) -> pd.Series:
    return _norm(keys).map(lambda k: mapping.get(str(k).strip(), "") if pd.notna(k) else "")


def _format_date_value(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "nat", "nat"}:
        return ""
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=False)
    if pd.isna(parsed):
        parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        # already dd-mm-yyyy?
        if len(text) >= 10 and text[2] == "-" and text[5] == "-":
            return text[:10]
        return text
    return parsed.strftime("%d-%b-%Y")


def _format_date_columns(df: pd.DataFrame, columns: set[str]) -> pd.DataFrame:
    out = df.copy()
    for col in columns:
        if col not in out.columns:
            continue
        out[col] = out[col].map(_format_date_value)
    return out


def _to_excel_number(value: object) -> float | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text.lower() in {"nan", "none", "nat", "<na>"}:
        return None
    num = pd.to_numeric(text, errors="coerce")
    if pd.isna(num):
        return None
    return float(num)


def _sanitize_result_for_excel(
    df: pd.DataFrame,
    numeric_columns: set[str],
) -> pd.DataFrame:
    """Keep numeric columns as numbers; sanitize other columns as Excel-safe text."""
    out = df.copy()
    for col in out.columns:
        if col in numeric_columns:
            out[col] = out[col].map(_to_excel_number)
        else:
            out[col] = out[col].map(_excel_safe_text)
    return out


def resolve_guide_path(project_root: Path) -> Path:
    for folder in GUIDE_DIR_CANDIDATES:
        path = project_root / folder / GUIDE_FILENAME
        if path.exists():
            return path
    tried = ", ".join(str(project_root / f / GUIDE_FILENAME) for f in GUIDE_DIR_CANDIDATES)
    raise ValueError(f"File guide tidak ditemukan. Dicari: {tried}")


def filter_rejection_blank(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """Keep rows where Reason for rejection is empty (Guide filter)."""
    rejection_col = find_column(
        df,
        [
            "Reason for rejection",
            "Reason For Rejection",
            "Rejection Reason",
        ],
    )
    rejection_values = df[rejection_col]
    rejection_blank = rejection_values.isna() | (
        rejection_values.astype(str).str.strip().isin(
            ["", "nan", "None", "NaT"]
        )
    )
    excluded = int((~rejection_blank).sum())
    return df.loc[rejection_blank].reset_index(drop=True), excluded


def filter_exclude_materials(
    df: pd.DataFrame,
    material_col_candidates: list[str] | None = None,
) -> tuple[pd.DataFrame, int]:
    candidates = material_col_candidates or [
        "Material",
        "Material No",
        "Material Number",
        "MATNR",
    ]
    material_col = find_column(df, candidates)
    values = _norm(df[material_col]).astype(str).str.upper()
    exclude = {m.upper() for m in EXCLUDE_MATERIAL_NOS}
    mask = values.isin(exclude)
    excluded = int(mask.sum())
    return df.loc[~mask].reset_index(drop=True), excluded


def filter_order_qty_ne_od_qty(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    order_col = find_column(
        df,
        ["Order Quantity", "Order Qty", "Ordered Quantity", "ORDER_QUANTITY"],
    )
    od_col = find_column(
        df,
        ["OD Quantity", "OD Qty", "Od Quantity", "OD_QUANTITY"],
    )
    order_qty = _to_numeric(df[order_col])
    od_qty = _to_numeric(df[od_col])
    equal = order_qty.notna() & od_qty.notna() & order_qty.eq(od_qty)
    excluded = int(equal.sum())
    return df.loc[~equal].reset_index(drop=True), excluded


def _od_group_key(df: pd.DataFrame) -> pd.Series:
    return (
        _norm(df["Sales document"]).fillna("").astype(str)
        + "|"
        + _norm(df["Sales Document Item"]).fillna("").astype(str)
        + "|"
        + _norm(df["Material No"]).fillna("").astype(str)
        + "|"
        + _qty_key(df["Order Quantity"])
    )


def _total_od_quantity(df: pd.DataFrame) -> pd.Series:
    """SUM(OD Quantity) per Sales document + Item + Material + Order Quantity."""
    od_qty = _to_numeric(df["OD Quantity"]).fillna(0)
    return od_qty.groupby(_od_group_key(df), sort=False).transform("sum")


def filter_order_qty_eq_total_od(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """
    Remove rows where Order Quantity equals SUM(OD Quantity) within group:
    Sales document + Sales Document Item + Material No + Order Quantity.
    """
    required = [*FULL_OD_GROUP_COLUMNS, "OD Quantity"]
    if any(col not in df.columns for col in required):
        return df.reset_index(drop=True), 0

    order_qty = _to_numeric(df["Order Quantity"])
    total_od = _total_od_quantity(df)
    full_od = order_qty.notna() & order_qty.eq(total_od)
    excluded = int(full_od.sum())
    return df.loc[~full_od].reset_index(drop=True), excluded


def filter_remove_duplicates(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """Drop exact duplicate rows on 7 key fields (keep first)."""
    keys = [c for c in DEDUPE_KEY_COLUMNS if c in df.columns]
    if not keys:
        return df.reset_index(drop=True), 0

    work = df.copy()
    for col in keys:
        if col in {"Order Quantity", "OD Quantity"}:
            work[col] = _to_numeric(work[col])
        else:
            work[col] = _norm(work[col]).fillna("")

    before = len(work)
    keep = ~work.duplicated(subset=keys, keep="first")
    excluded = before - int(keep.sum())
    return df.loc[keep].reset_index(drop=True), excluded


def _series_blank(series: pd.Series) -> pd.Series:
    return series.isna() | (
        series.astype(str)
        .str.strip()
        .isin(["", "nan", "None", "NaT", "<NA>"])
    )


def _as_date_only(series: pd.Series) -> pd.Series:
    """Parse values to midnight timestamps (date-only comparison)."""
    parsed = pd.to_datetime(series, errors="coerce", dayfirst=False)
    missing = parsed.isna()
    if missing.any():
        parsed2 = pd.to_datetime(series.loc[missing], errors="coerce", dayfirst=True)
        parsed = parsed.copy()
        parsed.loc[missing] = parsed2
    return parsed.dt.normalize()


def apply_action_rules(df: pd.DataFrame) -> pd.DataFrame:
    """
    Fill Action by Else-If first-match order (Guide business rules).
    """
    from datetime import date

    out = df.copy()
    n = len(out)
    empty = pd.Series([""] * n, index=out.index, dtype=object)

    storage_raw = out["Storage Location"] if "Storage Location" in out.columns else empty
    storage = _norm(storage_raw).fillna("").astype(str).str.upper()
    storage_blank = _series_blank(storage_raw)

    po_raw = (
        out["Purchasing Document"] if "Purchasing Document" in out.columns else empty
    )
    po_blank = _series_blank(po_raw)
    po_filled = ~po_blank

    klass = (
        _norm(out["Class"]).fillna("").astype(str).str.upper()
        if "Class" in out.columns
        else empty.astype(str)
    )

    milestone_raw = out["Milestone"] if "Milestone" in out.columns else empty
    milestone = _norm(milestone_raw).fillna("").astype(str)
    milestone_upper = milestone.str.upper()
    milestone_blank = _series_blank(milestone_raw)

    deletion_raw = (
        out["Deletion indicator"] if "Deletion indicator" in out.columns else empty
    )
    deletion_filled = ~_series_blank(deletion_raw)

    oq = (
        _to_numeric(out["Order Quantity"])
        if "Order Quantity" in out.columns
        else pd.Series([pd.NA] * n, index=out.index)
    )
    required_od_cols = [*FULL_OD_GROUP_COLUMNS, "OD Quantity"]
    if all(col in out.columns for col in required_od_cols):
        total_od = _total_od_quantity(out)
    else:
        total_od = (
            _to_numeric(out["OD Quantity"])
            if "OD Quantity" in out.columns
            else pd.Series([pd.NA] * n, index=out.index)
        )

    g38 = _to_numeric(out["1G38"]) if "1G38" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    s67 = _to_numeric(out["1S67"]) if "1S67" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    s66 = _to_numeric(out["1S66"]) if "1S66" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    s76 = _to_numeric(out["1S76"]) if "1S76" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    s81 = _to_numeric(out["1S81"]) if "1S81" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    sng = _to_numeric(out["SNG"]) if "SNG" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    mell = _to_numeric(out["Mell"]) if "Mell" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    qns = _to_numeric(out["QNS"]) if "QNS" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    sag = _to_numeric(out["SAG"]) if "SAG" in out.columns else pd.Series([pd.NA] * n, index=out.index)
    cpavail_sum = sng.fillna(0) + mell.fillna(0) + qns.fillna(0) + sag.fillna(0)

    vendor_raw = (
        out["Vendor/supplying plant"]
        if "Vendor/supplying plant" in out.columns
        else empty
    )
    vendor = _norm(vendor_raw).fillna("").astype(str).str.upper()
    vendor_cadc = vendor.str.startswith("1000085") | vendor.str.contains(
        "CATERPILLAR ASIA DELIVERY CENTER",
        regex=False,
        na=False,
    )

    agreement = (
        _norm(out["Agreement Type"]).fillna("").astype(str).str.upper()
        if "Agreement Type" in out.columns
        else empty.astype(str)
    )

    shipment_raw = (
        out["Shipment Number"] if "Shipment Number" in out.columns else empty
    )
    shipment_blank = _series_blank(shipment_raw)
    shipment_filled = ~shipment_blank

    if "PO Created Date" in out.columns:
        po_created = _as_date_only(out["PO Created Date"])
    else:
        po_created = pd.Series([pd.NaT] * n, index=out.index)
    if "Shp By Dt" in out.columns:
        shp_by = _as_date_only(out["Shp By Dt"])
    else:
        shp_by = pd.Series([pd.NaT] * n, index=out.index)
    today = pd.Timestamp(date.today()).normalize()
    shp_minus_today_days = (shp_by - today).dt.days

    bord = storage.eq("BORD")
    dmdv = storage.eq("DMDV")
    bord_or_dmdv = bord | dmdv

    action = empty.copy()

    def set_where(mask: pd.Series, text: str) -> None:
        nonlocal action
        apply_mask = mask.fillna(False) & action.eq("")
        action = action.mask(apply_mask, text)

    # Exact Else-If order from Guide business rules
    set_where(bord & po_blank, "Review & Submit BO")
    set_where(
        oq.notna() & total_od.notna() & oq.lt(total_od),
        "Cek Anomali (Possible double supply)",
    )
    set_where(milestone_upper.eq("GRIEFED"), "Griefed. Cek Antares EZ40")
    set_where(milestone_upper.eq("CANCELLED"), "Cancelled. Cek Antares EZ40")
    set_where(milestone_upper.eq("ESD NEEDED"), "Uplift to CPRO or Emergency")
    set_where(
        milestone_blank & po_created.notna() & po_created.eq(today),
        "Wait Transmit to CAT",
    )
    set_where(
        milestone_blank & po_created.notna() & po_created.lt(today),
        "Failed Transmit to CAT and Cek Antares EZ40",
    )
    set_where(
        bord & po_filled & g38.notna() & oq.notna() & g38.gt(oq),
        "BO Fill From Stock",
    )
    set_where(
        bord & po_filled & g38.notna() & g38.gt(0),
        "BO Fill From Stock Partial",
    )
    set_where(dmdv & po_blank & g38.notna() & g38.gt(0), "ReBO to stock")
    set_where(dmdv & po_blank & s67.notna() & s67.gt(0), "ReBO to 1S67")
    set_where(dmdv & po_blank & s66.notna() & s66.gt(0), "ReBO to 1S66")
    set_where(dmdv & po_blank & s76.notna() & s76.gt(0), "ReBO to 1S76")
    set_where(dmdv & po_blank & s81.notna() & s81.gt(0), "ReBO to 1S81")
    set_where(
        bord_or_dmdv & po_blank & oq.notna() & s67.notna() & oq.le(s67),
        "ReBO to 1S67",
    )
    set_where(
        bord_or_dmdv & po_blank & oq.notna() & s66.notna() & oq.le(s66),
        "ReBO to 1S66",
    )
    set_where(
        bord_or_dmdv & po_blank & oq.notna() & s76.notna() & oq.le(s76),
        "ReBO to 1S76",
    )
    set_where(
        bord_or_dmdv & po_blank & oq.notna() & s81.notna() & oq.le(s81),
        "ReBO to 1S81",
    )
    set_where(
        bord_or_dmdv & po_filled & oq.notna() & s67.notna() & oq.le(s67),
        "Request STO from 1S67",
    )
    set_where(
        bord_or_dmdv & po_filled & oq.notna() & s66.notna() & oq.le(s66),
        "Request STO from 1S66",
    )
    set_where(
        bord_or_dmdv & po_filled & oq.notna() & s76.notna() & oq.le(s76),
        "Request STO from 1S76",
    )
    set_where(
        bord_or_dmdv & po_filled & oq.notna() & s81.notna() & oq.le(s81),
        "Request STO from 1S81",
    )
    set_where(storage_blank, "Cek & Create OD & F/u GI")
    set_where(klass.eq("ON-ORDER"), "Cek On-Order")
    hubs_zero = (
        s67.fillna(0).eq(0)
        & s66.fillna(0).eq(0)
        & s76.fillna(0).eq(0)
        & s81.fillna(0).eq(0)
    )
    set_where(
        klass.eq("ON-HAND") & hubs_zero,
        "ReBO to CAT & Cek Availability Incountry",
    )
    set_where(deletion_filled, "PO Possible Delete / BO Cancelled")
    set_where(
        vendor_cadc & milestone_upper.eq("SHIPPED"),
        "Keep Monitor",
    )
    set_where(
        agreement.eq("CPRO")
        & milestone_upper.eq("SOURCED")
        & shp_minus_today_days.eq(27),
        "Request Early Invoice",
    )
    set_where(
        milestone_upper.eq("SOURCED") & sng.notna() & oq.notna() & sng.gt(oq),
        "Keep Monitor (Milestone Sourced)",
    )
    set_where(
        milestone_upper.eq("SOURCED") & sng.notna() & oq.notna() & sng.lt(oq),
        "F/u invoice (Milestone Sourced)",
    )
    set_where(
        milestone_upper.eq("ESD AVAILABLE")
        & oq.notna()
        & oq.le(cpavail_sum),
        "Fu/Keep Monitor",
    )
    set_where(
        klass.eq("TRANSFER") & shipment_filled,
        "Keep Monitor (BO Incountry)",
    )
    set_where(klass.eq("TRANSFER") & shipment_blank, "F/u BR")
    set_where(
        milestone_upper.eq("SHIPPED"),
        "Keep Monitor (Milestone Shipped)",
    )
    set_where(
        milestone_upper.eq("ESD AVAILABLE"),
        "Keep Monitor (Milestone ESD Available)",
    )
    set_where(
        milestone_blank & po_filled,
        "Possible Failed Transmit to CAT and Cek Antares EZ40",
    )
    set_where(
        milestone_upper.eq("SOURCED") & po_filled,
        "F/u invoice (Milestone Sourced)",
    )
    # Else → leave blank

    out["Action"] = action
    return out


def _format_estimasi_date(base_dates: pd.Series, days: pd.Series) -> pd.Series:
    """Add day offsets to dates; return dd-mmm-yyyy text (blank if base missing)."""
    parsed = _as_date_only(base_dates)
    offset = _to_numeric(days).fillna(0)
    out: list[str] = []
    for dt, day in zip(parsed.tolist(), offset.tolist(), strict=False):
        if pd.isna(dt):
            out.append("")
            continue
        try:
            day_n = int(float(day)) if pd.notna(day) else 0
        except (TypeError, ValueError):
            day_n = 0
        out.append((dt + pd.Timedelta(days=day_n)).strftime("%d-%b-%Y"))
    return pd.Series(out, index=base_dates.index, dtype=object)


def apply_estimasi_remark_rules(df: pd.DataFrame) -> pd.DataFrame:
    """
    Fill Estimasi + Remark from sheet 'Estimasi & Remark' (Else-If first-match).
    More specific rules first, then general Purchasing Document blank.
    Template column is 'Remark' (legacy 'Remaks' still accepted).
    """
    from datetime import date

    out = df.copy()
    n = len(out)
    empty = pd.Series([""] * n, index=out.index, dtype=object)
    estimasi = empty.copy()
    remark = empty.copy()
    filled = pd.Series([False] * n, index=out.index)

    def set_where(mask: pd.Series, est: pd.Series | str, rem: str) -> None:
        nonlocal estimasi, remark, filled
        apply = mask.fillna(False) & ~filled
        if not apply.any():
            return
        if isinstance(est, str):
            estimasi = estimasi.where(~apply, est)
        else:
            estimasi = estimasi.where(~apply, est)
        remark = remark.where(~apply, rem)
        filled = filled | apply

    po_raw = (
        out["Purchasing Document"] if "Purchasing Document" in out.columns else empty
    )
    po_blank = _series_blank(po_raw)
    po_filled = ~po_blank

    klass = (
        _norm(out["Class"]).fillna("").astype(str).str.upper()
        if "Class" in out.columns
        else empty.astype(str)
    )
    shipment_raw = (
        out["Shipment Number"] if "Shipment Number" in out.columns else empty
    )
    shipment_blank = _series_blank(shipment_raw)
    shipment_filled = ~shipment_blank

    vendor_raw = (
        out["Vendor/supplying plant"]
        if "Vendor/supplying plant" in out.columns
        else empty
    )
    vendor = _norm(vendor_raw).fillna("").astype(str).str.upper()
    vendor_cadc = vendor.str.startswith("1000085") | vendor.str.contains(
        "CATERPILLAR ASIA DELIVERY CENTER",
        regex=False,
        na=False,
    )

    agreement = (
        _norm(out["Agreement Type"]).fillna("").astype(str).str.upper()
        if "Agreement Type" in out.columns
        else empty.astype(str)
    )
    agreement_cpro = agreement.eq("CPRO")
    agreement_mega_down = agreement.str.contains("MEGA", na=False) | agreement.str.contains(
        "DOWN",
        na=False,
    )
    agreement_blank = _series_blank(
        out["Agreement Type"] if "Agreement Type" in out.columns else empty
    )

    milestone_raw = out["Milestone"] if "Milestone" in out.columns else empty
    milestone_upper = _norm(milestone_raw).fillna("").astype(str).str.upper()

    eta = out["ETA D"] if "ETA D" in out.columns else empty
    today = pd.Series(
        [pd.Timestamp(date.today())] * n,
        index=out.index,
    )
    today_plus_eta = _format_estimasi_date(today, eta)
    today_plus_120 = _format_estimasi_date(
        today,
        pd.Series([120] * n, index=out.index),
    )

    ship_num_date = (
        out["Shipment Number Date"]
        if "Shipment Number Date" in out.columns
        else empty
    )
    shp_by = out["Shp By Dt"] if "Shp By Dt" in out.columns else empty
    act_dept = (
        out["Act Dept Dt -Inv CAT"]
        if "Act Dept Dt -Inv CAT" in out.columns
        else empty
    )
    est_ship = out["Est Ship Date"] if "Est Ship Date" in out.columns else empty
    source_date = out["Source Date"] if "Source Date" in out.columns else empty

    # SOURCE DATE + 3 + ETA
    source_plus_3_eta = _format_estimasi_date(
        source_date,
        _to_numeric(eta).fillna(0) + 3,
    )

    # 1) CLASS = ASSY + PO blank
    set_where(
        klass.eq("ASSY") & po_blank,
        "TBA",
        "PO Subcont BUILD UP",
    )
    # 2) CLASS = TRANSFER + PO filled + Shipment blank → TODAY()+ETA
    set_where(
        klass.eq("TRANSFER") & po_filled & shipment_blank,
        today_plus_eta,
        "Follow Up Cabang",
    )
    # 3) CLASS = TRANSFER + Shipment filled → Shipment Number Date + ETA
    set_where(
        klass.eq("TRANSFER") & shipment_filled,
        _format_estimasi_date(ship_num_date, eta),
        "Follow Up CKB",
    )

    po_cadc = klass.eq("PO") & vendor_cadc

    # CPRO milestones
    set_where(
        po_cadc & agreement_cpro & milestone_upper.eq("SOURCED"),
        _format_estimasi_date(shp_by, eta),
        "Waiting Inv",
    )
    set_where(
        po_cadc & agreement_cpro & milestone_upper.eq("SHIPPED"),
        _format_estimasi_date(act_dept, eta),
        "Follow Up CKB",
    )
    set_where(
        po_cadc & agreement_cpro & milestone_upper.eq("FUTURE DATED"),
        _format_estimasi_date(shp_by, eta),
        "F/U ESD",
    )
    set_where(
        po_cadc & agreement_cpro & milestone_upper.eq("ESD NEEDED"),
        today_plus_120,
        "F/U ESD",
    )
    set_where(
        po_cadc & agreement_cpro & milestone_upper.eq("ESD AVAILABLE"),
        _format_estimasi_date(est_ship, eta),
        "F/U ESD",
    )

    # MEGA / DOWN
    set_where(
        po_cadc & agreement_mega_down & milestone_upper.eq("ESD NEEDED"),
        "TBA",
        "Waiting for branch release or Caterpillar submit",
    )
    set_where(
        po_cadc & agreement_mega_down & milestone_upper.eq("ESD AVAILABLE"),
        _format_estimasi_date(est_ship, eta),
        "F/U ESD",
    )

    # Agreement blank
    set_where(
        po_cadc & agreement_blank & _series_blank(milestone_raw),
        "TBA",
        "Waiting for branch release or Caterpillar submit",
    )
    set_where(
        po_cadc & agreement_blank & milestone_upper.eq("SOURCED"),
        source_plus_3_eta,
        "Waiting invoice",
    )
    set_where(
        po_cadc & agreement_blank & milestone_upper.eq("SHIPPED"),
        _format_estimasi_date(act_dept, eta),
        "Waiting handover to CKB",
    )
    set_where(
        po_cadc & agreement_blank & milestone_upper.eq("ESD NEEDED"),
        today_plus_120,
        "F/U ESD",
    )
    set_where(
        po_cadc & agreement_blank & milestone_upper.eq("ESD AVAILABLE"),
        _format_estimasi_date(est_ship, eta),
        "F/U ESD",
    )

    # General: Purchasing Document blank
    set_where(
        po_blank,
        "TBA",
        "Waiting for branch release or Caterpillar submit",
    )

    out["Estimasi"] = estimasi
    # Data Template header is now "Remark" (was "Remaks")
    out["Remark"] = remark
    if "Remaks" in out.columns:
        out["Remaks"] = remark
    return out


def _rename_from_source(source: pd.DataFrame) -> pd.DataFrame:
    """Map Source Item columns onto Data Template field names."""
    mapping = [
        ("Customer Reference Number", ["Customer Reference Number"]),
        ("Plant", ["Plant"]),
        ("Customer Name", ["Customer Name"]),
        ("SO Date", ["Created on", "SO Date"]),
        ("Customer PO", ["PO Number", "Customer PO", "Purchase order no."]),
        ("IREQ Item", ["IREQ Item"]),
        ("Sales document", ["Sales Document", "Sales document"]),
        ("Sales Document Item", ["Item (SD)", "Sales Document Item"]),
        # Guide: kolom berisi PO/TRANSFER/ON-HAND → Class.1 di export SAP
        ("Class", ["Class.1", "Class", "Class.2"]),
        ("Created by", ["Created by"]),
        ("Material No", ["Material", "Material No", "Material Number"]),
        ("Material Description", ["Material Description", "Description"]),
        ("Order Quantity", ["Order Quantity"]),
        ("Purchase Requisition", ["Purchase Requisition"]),
        ("PR Item", ["Item of Requisition", "PR Item"]),
        ("Purchasing Document", ["Purchasing Document"]),
        ("PO Item", ["Item", "PO Item"]),
        ("OD of Sales Order Item", ["OD of Sales Order Item"]),
        ("OD Item of Sales Order Item", ["OD Item of Sales Order Item"]),
        ("Act GI Date of OD", ["Act GI Date of OD"]),
        ("OD Quantity", ["OD Quantity"]),
        ("Gate Pass", ["Gate Pass"]),
        ("Gate Pass Print Date", ["Gate Pass Print Date"]),
        ("Storage Location", ["Storage Location"]),
        (
            "Reason for rejection",
            [
                "Reason for rejection",
                "Reason For Rejection",
                "Rejection Reason",
            ],
        ),
        (
            "Deletion indicator",
            [
                "Deletion indicator",
                "Deletion Indicator",
                "Deletion Ind.",
                "Del. Ind.",
            ],
        ),
    ]
    out = pd.DataFrame(index=source.index)
    for target, candidates in mapping:
        col = _safe_col(source, candidates)
        out[target] = _norm(source[col]) if col else ""
    return out


@dataclass
class BoGuideResult:
    job_id: str
    row_count: int
    excluded_qty_equal_count: int
    excluded_delivery_charge_count: int
    excluded_rejection_count: int
    excluded_exact_duplicate_count: int
    excluded_duplicate_count: int
    source_item_files: list[str]
    guide_path: str
    files: dict[str, Path]
    preview: list[dict[str, str]]
    missing_sources: list[str]


def _excel_safe_text(value: object) -> str:
    """Strip characters that Excel XML rejects (causes 'Repaired' on open)."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value)
    if text.lower() in {"nan", "none", "nat", "<na>"}:
        return ""
    # XML 1.0 illegal control chars
    return "".join(
        ch for ch in text if ord(ch) >= 32 or ch in "\t\n\r"
    )


def _excel_cell_value(cell_value: object) -> object | None:
    if cell_value is None or (isinstance(cell_value, float) and pd.isna(cell_value)):
        return None
    if hasattr(cell_value, "item") and not isinstance(cell_value, str):
        try:
            cell_value = cell_value.item()
        except Exception:  # noqa: BLE001
            pass
    if isinstance(cell_value, float) and pd.isna(cell_value):
        return None
    return cell_value


def _write_sheet_dataframe(
    ws,
    df: pd.DataFrame,
    headers: list[str],
) -> None:
    """Clear body and write dataframe rows aligned to headers (row 1 kept/set)."""
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for col_idx, label in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=label)
    if df.empty:
        return
    for row_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_idx, cell_value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_excel_cell_value(cell_value))


def filter_template_no_po(df: pd.DataFrame) -> pd.DataFrame:
    """Purchasing Document blank AND Storage Location = BORD."""
    if df.empty:
        return df.copy()
    empty = pd.Series([""] * len(df), index=df.index, dtype=object)
    po_raw = df["Purchasing Document"] if "Purchasing Document" in df.columns else empty
    sloc_raw = df["Storage Location"] if "Storage Location" in df.columns else empty
    po_blank = _series_blank(po_raw)
    bord = _norm(sloc_raw).fillna("").astype(str).str.upper().eq("BORD")
    return df.loc[po_blank & bord].reset_index(drop=True)


def build_template_no_po_frame(
    result_df: pd.DataFrame,
    no_po_headers: list[str],
) -> pd.DataFrame:
    """Map Data Template columns onto Template NO PO layout."""
    out = pd.DataFrame(index=result_df.index)
    # Prefer first occurrence when duplicate column labels exist
    col_positions: dict[str, int] = {}
    for i, c in enumerate(result_df.columns):
        key = str(c).strip()
        if key not in col_positions:
            col_positions[key] = i
    lower_positions: dict[str, int] = {}
    for key, i in col_positions.items():
        lower_positions.setdefault(key.lower(), i)

    for header in no_po_headers:
        label = str(header).strip() if header is not None else ""
        aliases = NO_PO_COLUMN_ALIASES.get(label, (label,))
        pos = None
        for alias in aliases:
            if alias in col_positions:
                pos = col_positions[alias]
                break
            low = alias.lower()
            if low in lower_positions:
                pos = lower_positions[low]
                break
        if pos is None:
            out[label] = ""
        else:
            series = result_df.iloc[:, pos]
            if isinstance(series, pd.DataFrame):
                series = series.iloc[:, 0]
            out[label] = series
    return out.reindex(
        columns=[str(h).strip() if h is not None else "" for h in no_po_headers]
    )


def _write_bo_report_workbook(
    guide_path: Path,
    output_path: Path,
    result_df: pd.DataFrame,
    template_headers: list[str],
    output_headers: list[str],
    no_po_df: pd.DataFrame | None = None,
    no_po_headers: list[str] | None = None,
) -> None:
    """
    Write by copying the guide workbook (keeps Guide sheet + valid OOXML),
    then fill Data Template and Template NO PO.
    """
    from openpyxl import load_workbook

    wb = load_workbook(guide_path)
    if TEMPLATE_SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet '{TEMPLATE_SHEET}' tidak ada di guide.")
    ws = wb[TEMPLATE_SHEET]

    _write_sheet_dataframe(
        ws,
        result_df.reindex(columns=template_headers),
        output_headers,
    )

    # Template NO PO: always derive from result_df at write time
    if TEMPLATE_NO_PO_SHEET in wb.sheetnames:
        ws_no = wb[TEMPLATE_NO_PO_SHEET]
        headers = [str(h).strip() for h in (no_po_headers or []) if h]
        if not headers:
            headers = [
                str(cell.value).strip()
                for cell in next(ws_no.iter_rows(min_row=1, max_row=1))
                if cell.value is not None
            ]
        if not headers:
            headers = list(NO_PO_COLUMN_ALIASES.keys())
        frame = build_template_no_po_frame(filter_template_no_po(result_df), headers)
        _write_sheet_dataframe(ws_no, frame.reindex(columns=headers), headers)

    # Keep guide sheets (Data Template, Guide, Estimasi & Remark, Template NO PO)
    for name in list(wb.sheetnames):
        if name not in GUIDE_KEEP_SHEETS:
            del wb[name]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()


def _unique_headers(headers: list[object]) -> list[str]:
    """Make DataFrame-safe unique names; second Plant -> Plant.1."""
    seen: dict[str, int] = {}
    unique: list[str] = []
    for raw in headers:
        name = str(raw).strip() if raw is not None else ""
        if name not in seen:
            seen[name] = 0
            unique.append(name)
        else:
            seen[name] += 1
            unique.append(f"{name}.{seen[name]}")
    return unique


def generate_bo_report(
    project_root: Path,
    output_dir: Path,
) -> BoGuideResult:
    guide_path = resolve_guide_path(project_root)

    raw_headers = list(
        pd.read_excel(guide_path, sheet_name=TEMPLATE_SHEET, nrows=0).columns
    )
    # Prefer openpyxl for true duplicate header labels (two Plant columns)
    try:
        from openpyxl import load_workbook

        wb = load_workbook(guide_path, read_only=True, data_only=True)
        ws = wb[TEMPLATE_SHEET]
        excel_headers = [
            cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))
        ]
        wb.close()
        if excel_headers and any(excel_headers):
            raw_headers = excel_headers
    except Exception:  # noqa: BLE001
        pass

    template_headers = _unique_headers(raw_headers)
    output_headers = [str(h).strip() if h is not None else "" for h in raw_headers]
    header_po_qty = next(
        (h for h in template_headers if str(h).strip() == "PO Quantity"),
        "PO Quantity",
    )
    plant_like = [h for h in template_headers if str(h).strip().startswith("Plant")]
    header_plant_formula = next(
        (h for h in plant_like if h != "Plant"),
        "Plant.1",
    )

    source_dir = project_root / "data-sap-zvsd_parts_progress-source_item"
    source_files = [p.name for p in list_excel_files(source_dir)]
    source_raw = _load_folder_frames(source_dir)
    if source_raw.empty:
        raise ValueError("Folder data-sap-zvsd_parts_progress-source_item kosong.")

    source_filtered, excluded_qty_equal_count = filter_order_qty_ne_od_qty(source_raw)
    source_filtered, excluded_delivery_charge_count = filter_exclude_materials(
        source_filtered
    )
    source_filtered, excluded_rejection_count = filter_rejection_blank(source_filtered)
    base = _rename_from_source(source_filtered)

    missing_sources: list[str] = []

    # --- Parts Progress (by Sales document) ---
    pp_dir = project_root / "data-sap-zvsd_parts_progress"
    pp = _load_folder_frames(pp_dir)
    if pp.empty:
        missing_sources.append("data-sap-zvsd_parts_progress")
        for col in ("Released Date", "Need By Date", "Ordered By"):
            base[col] = ""
    else:
        sd_col = find_column(pp, ["Sales Document", "Sales document"])
        released = _safe_col(pp, ["Released Date"])
        need_by = _safe_col(pp, ["Need by Date", "Need By Date"])
        ordered_by = _safe_col(pp, ["Salesman Name", "Ordered By"])
        key = _norm(pp[sd_col])
        base["Released Date"] = _map_series(
            base["Sales document"],
            _lookup_map(pp, key, released) if released else {},
        )
        base["Need By Date"] = _map_series(
            base["Sales document"],
            _lookup_map(pp, key, need_by) if need_by else {},
        )
        base["Ordered By"] = _map_series(
            base["Sales document"],
            _lookup_map(pp, key, ordered_by) if ordered_by else {},
        )

    # --- Manuf = RIGHT(Material No, 2) ---
    base["Manuf"] = _norm(base["Material No"]).astype(str).str[-2:]
    base.loc[_norm(base["Material No"]).isna(), "Manuf"] = ""

    # --- PO Moni lookups (Guide: different keys per field) ---
    po_dir = project_root / "data-sap-zmpu_po_moni"
    po = _load_folder_frames(po_dir)
    # Key = Purchasing Document only
    po_by_pd = {
        "PO Created Date": ["Document Date", "PO Created Date"],
        "Order Class": ["Order Class"],
        "Vendor/supplying plant": ["Vendor/supplying plant"],
    }
    # Key = SO + Item + Material + Purchasing Document (no Order Qty)
    po_by_so_item_mat_pd = {
        header_po_qty: ["Order Quantity", "PO Quantity"],
    }
    # Key = SO + Item + Material + Order Qty + Purchasing Document
    po_by_full = {
        "Agreement Type": ["Agreement Type"],
        "Ship Out": ["Ship Out"],
        "Ship In": ["Ship In"],
        "Invoice Date": ["Invoice Date"],
        "Invoice Reference Number": ["Invoice Reference Number"],
        "Shipment Number": ["Shipment Number"],
        "Shipment Number Date": ["Shipment Number Date"],
    }
    all_po_targets = [
        *po_by_pd.keys(),
        *po_by_so_item_mat_pd.keys(),
        *po_by_full.keys(),
    ]
    if po.empty:
        missing_sources.append("data-sap-zmpu_po_moni")
        for col in all_po_targets:
            base[col] = ""
    else:
        so_c = find_column(po, ["Sales Order Number", "Sales Document", "Sales document"])
        item_c = find_column(po, ["Sales Order Item", "Item (SD)", "Sales Document Item"])
        mat_c = find_column(po, ["Material", "Material No", "Material Number"])
        qty_c = find_column(po, ["Order Quantity"])
        pd_c = find_column(po, ["Purchasing Document"])

        po_pd_key = _norm(po[pd_c]).fillna("")
        base_pd_key = _norm(base["Purchasing Document"]).fillna("")

        po_so_item_mat_pd = (
            _norm(po[so_c]).fillna("")
            + "|"
            + _norm(po[item_c]).fillna("")
            + "|"
            + _norm(po[mat_c]).fillna("")
            + "|"
            + _norm(po[pd_c]).fillna("")
        )
        base_so_item_mat_pd = (
            _norm(base["Sales document"]).fillna("")
            + "|"
            + _norm(base["Sales Document Item"]).fillna("")
            + "|"
            + _norm(base["Material No"]).fillna("")
            + "|"
            + _norm(base["Purchasing Document"]).fillna("")
        )

        po_full_key = (
            _norm(po[so_c]).fillna("")
            + "|"
            + _norm(po[item_c]).fillna("")
            + "|"
            + _norm(po[mat_c]).fillna("")
            + "|"
            + _qty_key(po[qty_c])
            + "|"
            + _norm(po[pd_c]).fillna("")
        )
        base_full_key = (
            _norm(base["Sales document"]).fillna("")
            + "|"
            + _norm(base["Sales Document Item"]).fillna("")
            + "|"
            + _norm(base["Material No"]).fillna("")
            + "|"
            + _qty_key(base["Order Quantity"])
            + "|"
            + _norm(base["Purchasing Document"]).fillna("")
        )

        def _apply_po_maps(
            field_map: dict[str, list[str]],
            source_keys: pd.Series,
            target_keys: pd.Series,
        ) -> None:
            for target, candidates in field_map.items():
                value_col = _safe_col(po, candidates)
                mapping: dict[str, str] = {}
                if value_col:
                    for k, v in zip(
                        source_keys.tolist(),
                        _norm(po[value_col]).tolist(),
                        strict=False,
                    ):
                        if k and k not in mapping:
                            mapping[k] = v
                base[target] = target_keys.map(lambda k, m=mapping: m.get(k, ""))

        _apply_po_maps(po_by_pd, po_pd_key, base_pd_key)
        _apply_po_maps(po_by_so_item_mat_pd, po_so_item_mat_pd, base_so_item_mat_pd)
        _apply_po_maps(po_by_full, po_full_key, base_full_key)

    # Plant.1 = LEFT(Vendor/supplying plant, 4)
    vendor = _norm(base.get("Vendor/supplying plant", pd.Series([""] * len(base))))
    base[header_plant_formula] = vendor.astype(str).str[:4]
    base.loc[vendor.isna() | (vendor == ""), header_plant_formula] = ""

    # --- CPAvail by Material (base, strip :suffix) ---
    cp_dir = project_root / "data-sap-zmim_cpavail"
    cp = _load_folder_frames(cp_dir)
    cp_map_fields = {
        "SNG": ["SNG (Singapore)", "SNG"],
        "Mell": ["MEL (Melbourne)", "Mell", "MEL"],
        "QNS": ["QNS (Queensland)", "QNS"],
        "SAG": ["SAG (Sagami)", "SAG"],
    }
    base_mat_base = _material_base(base["Material No"])
    if cp.empty:
        missing_sources.append("data-sap-zmim_cpavail")
        for col in cp_map_fields:
            base[col] = ""
    else:
        mat_c = find_column(cp, ["Material", "Material Number", "Material No"])
        cp_mat_base = _material_base(cp[mat_c])
        for target, candidates in cp_map_fields.items():
            value_col = _safe_col(cp, candidates)
            if value_col:
                base[target] = _map_by_material_base(
                    base_mat_base,
                    cp_mat_base,
                    _norm(cp[value_col]),
                )
            else:
                base[target] = ""

    # --- Stock info Order Method + plant 1G38 fields (Guide + Data Template) ---
    stock_dir = project_root / "data-sap-zmmm_stock_info"
    stock = _load_folder_frames(stock_dir)
    stock_template_names = [
        name
        for aliases, _src in STOCK_1G38_FIELD_SPECS
        for name in _template_targets(template_headers, aliases)
    ]
    stock_targets = ["Order Method", "1G38", *stock_template_names]
    if stock.empty:
        missing_sources.append("data-sap-zmmm_stock_info")
        for col in stock_targets:
            base[col] = ""
    else:
        mat_c = find_column(stock, ["Material Number", "Material", "Material No"])
        method_c = _safe_col(stock, ["Order Method"])
        plant_c = _safe_col(stock, ["Plant"])
        # Guide: 1G38 = On-hand Stock (ATP)
        avail_c = _safe_col(stock, list(STOCK_QTY_CANDIDATES))
        stock_mat_base = _material_base(stock[mat_c])

        if method_c:
            base["Order Method"] = _map_by_material_base(
                base_mat_base,
                stock_mat_base,
                _norm(stock[method_c]),
            )
        else:
            base["Order Method"] = ""

        stock_1g38 = (
            stock.loc[_norm(stock[plant_c]) == "1G38"]
            if plant_c
            else stock.iloc[0:0]
        )
        stock_1g38_mat = (
            _material_base(stock_1g38[mat_c])
            if not stock_1g38.empty
            else pd.Series(dtype=object)
        )
        if avail_c is not None and not stock_1g38.empty:
            base["1G38"] = _map_by_material_base(
                base_mat_base,
                stock_1g38_mat,
                _norm(stock_1g38[avail_c]),
            )
        else:
            base["1G38"] = ""

        for aliases, source_cands in STOCK_1G38_FIELD_SPECS:
            value_col = _safe_col(stock, list(source_cands))
            mapped = (
                _map_by_material_base(
                    base_mat_base,
                    stock_1g38_mat,
                    _norm(stock_1g38[value_col]),
                )
                if value_col is not None and not stock_1g38.empty
                else pd.Series([""] * len(base), index=base.index)
            )
            for target in _template_targets(template_headers, aliases):
                base[target] = mapped

    # --- PartViz lookups (material base + cust ref) ---
    partviz_path = project_root / "output" / "partviz_merged.xlsx"
    partviz_fields = {
        "Prim PSO": ["Prim PSO"],
        "Milestone": ["Milestone"],
        "Old ESD": ["Prv Est Ship Date", "Old ESD"],
        "Est Ship Date": ["Est Ship Date"],
        "ESD Flag": ["ESD Flag"],
        "Act Dept Dt -Inv CAT": ["Act Dept Dt", "Act Dept Dt -Inv CAT"],
        "Shp By Dt": ["Shp By Dt"],
        "Source Date": ["Source Date/Time", "Source Date"],
    }
    if not partviz_path.exists():
        missing_sources.append("output/partviz_merged.xlsx")
        for col in partviz_fields:
            base[col] = ""
    else:
        pv = read_excel_source(partviz_path)
        part_c = find_column(pv, ["Part No", "Material", "Material No"])
        cust_c = find_column(pv, ["Cust Ref", "Customer Reference Number"])
        pv_key = _material_base(pv[part_c]).fillna("") + "|" + _norm(pv[cust_c]).fillna("")
        base_pv_key = (
            _material_base(base["Material No"]).fillna("")
            + "|"
            + _norm(base["Customer Reference Number"]).fillna("")
        )
        for target, candidates in partviz_fields.items():
            value_col = _safe_col(pv, candidates)
            mapping: dict[str, str] = {}
            if value_col:
                for k, v in zip(pv_key.tolist(), _norm(pv[value_col]).tolist(), strict=False):
                    if k and k not in mapping:
                        mapping[k] = v
            base[target] = base_pv_key.map(lambda k: mapping.get(k, ""))

    # --- Total Price = Price per Material * Order Quantity ---
    price_path = project_root / "output" / "order_item_price_per_material.xlsx"
    if not price_path.exists():
        missing_sources.append("output/order_item_price_per_material.xlsx")
        base["Total Price"] = ""
    else:
        price_df = read_excel_source(price_path)
        mat_c = find_column(price_df, ["Material", "Material No", "Material Number"])
        so_c = find_column(price_df, ["Sales Document", "Sales document"])
        price_c = find_column(price_df, ["Price per Material"])
        price_key = _norm(price_df[mat_c]).fillna("") + "|" + _norm(price_df[so_c]).fillna("")
        mapping: dict[str, float] = {}
        prices = _to_numeric(price_df[price_c])
        for k, v in zip(price_key.tolist(), prices.tolist(), strict=False):
            if k and k not in mapping and pd.notna(v):
                mapping[k] = float(v)
        base_price_key = (
            _norm(base["Material No"]).fillna("")
            + "|"
            + _norm(base["Sales document"]).fillna("")
        )
        unit = base_price_key.map(lambda k: mapping.get(k, float("nan")))
        qty = _to_numeric(base["Order Quantity"])
        total = pd.Series(unit, index=base.index) * qty
        base["Total Price"] = total.apply(
            lambda v: "" if pd.isna(v) else f"{v:.4f}".rstrip("0").rstrip(".")
        )

    # --- RMS ETA = Plant.1 & Ship In (fallback Plant) ---
    plant_for_eta = base[header_plant_formula].where(
        _norm(base[header_plant_formula]).fillna("") != "",
        base["Plant"],
    )
    ship_in = _norm(base.get("Ship In", pd.Series([""] * len(base)))).fillna("")
    base["RMS ETA"] = plant_for_eta.fillna("").astype(str) + ship_in.astype(str)
    base.loc[(plant_for_eta.fillna("") == "") | (ship_in == ""), "RMS ETA"] = ""

    # --- Estimasi ETA D from data-estimasi ---
    est_dir = project_root / "data-estimasi"
    est = _load_folder_frames(est_dir)
    if est.empty:
        missing_sources.append("data-estimasi")
        base["ETA D"] = ""
    else:
        key_c = find_column(est, ["PlntShip", "PlantShip"])
        val_c = find_column(est, ["Est (Day)", "Est Day", "Est"])
        mapping = {}
        for k, v in zip(_norm(est[key_c]).tolist(), _norm(est[val_c]).tolist(), strict=False):
            if k and k not in mapping:
                mapping[k] = v
        base["ETA D"] = _map_series(base["RMS ETA"], mapping)

    # --- Hub stock by Material base + Plant (Guide: On-hand Stock ATP) ---
    hub_dir = project_root / "data-sap-zmmm_stock_info-hub"
    hub = _load_folder_frames(hub_dir)
    hub_plants = ["1S67", "1S66", "1S76", "1S81"]
    if hub.empty:
        missing_sources.append("data-sap-zmmm_stock_info-hub")
        for plant in hub_plants:
            base[plant] = ""
    else:
        mat_c = find_column(hub, ["Material Number", "Material", "Material No"])
        plant_c = find_column(hub, ["Plant"])
        value_c = _safe_col(hub, list(STOCK_QTY_CANDIDATES))
        hub_mat_base = _material_base(hub[mat_c])
        for plant in hub_plants:
            if value_c:
                mask = _norm(hub[plant_c]) == plant
                sub_keys = hub_mat_base.loc[mask]
                sub_vals = _norm(hub.loc[mask, value_c])
                base[plant] = _map_by_material_base(base_mat_base, sub_keys, sub_vals)
            else:
                base[plant] = ""

    # --- BO last Estimasi before / Remaks before ---
    bo_dir = project_root / "data-bo-last"
    bo_last = _load_folder_frames(bo_dir)
    if bo_last.empty:
        missing_sources.append("data-bo-last")
        base["Estimasi before"] = ""
        base["Remaks before"] = ""
    else:
        key_c = find_column(bo_last, ["Key (SO+PO+PN+QTY)", "Key"])
        est_c = _safe_col(bo_last, ["Estimasi"])
        rem_c = _safe_col(bo_last, ["Remaks", "Remarks", "Remark"])
        # Build same key style: SO + PO + Material + Qty
        base_bo_key = (
            _norm(base["Sales document"]).fillna("")
            + _norm(base["Purchasing Document"]).fillna("")
            + _norm(base["Material No"]).fillna("")
            + _qty_key(base["Order Quantity"])
        )
        for target, value_col in (("Estimasi before", est_c), ("Remaks before", rem_c)):
            mapping: dict[str, str] = {}
            if value_col:
                for k, v in zip(
                    _norm(bo_last[key_c]).tolist(),
                    _norm(bo_last[value_col]).tolist(),
                    strict=False,
                ):
                    if k and k not in mapping:
                        mapping[k] = v
            base[target] = base_bo_key.map(lambda k: mapping.get(k, ""))

    # Empty guide fields (filled later by rules)
    if "Estimasi" not in base.columns:
        base["Estimasi"] = ""
    if "Remark" not in base.columns:
        base["Remark"] = ""
    if "Remaks" not in base.columns and "Remaks" in template_headers:
        base["Remaks"] = ""
    if "Action" not in base.columns:
        base["Action"] = ""

    # Align to template headers (unique internal names)
    for col in template_headers:
        if col not in base.columns:
            base[col] = ""
    result_df = base.reindex(columns=template_headers)
    result_df = _format_date_columns(result_df, DATE_OUTPUT_COLUMNS)
    result_df, excluded_exact_duplicate_count = filter_remove_duplicates(result_df)
    result_df, excluded_duplicate_count = filter_order_qty_eq_total_od(result_df)
    result_df = apply_action_rules(result_df)
    result_df = apply_estimasi_remark_rules(result_df)
    result_df = _sanitize_result_for_excel(result_df, NUMERIC_OUTPUT_COLUMNS)

    # Template NO PO headers + filtered rows (PO blank + SLOC BORD)
    no_po_headers: list[str] = []
    try:
        from openpyxl import load_workbook

        wb_no = load_workbook(guide_path, read_only=True, data_only=True)
        if TEMPLATE_NO_PO_SHEET in wb_no.sheetnames:
            ws_no = wb_no[TEMPLATE_NO_PO_SHEET]
            no_po_headers = [
                str(cell.value).strip()
                for cell in next(ws_no.iter_rows(min_row=1, max_row=1))
                if cell.value is not None
            ]
        wb_no.close()
    except Exception:  # noqa: BLE001
        no_po_headers = list(NO_PO_COLUMN_ALIASES.keys())
    if not no_po_headers:
        no_po_headers = list(NO_PO_COLUMN_ALIASES.keys())
    no_po_df = build_template_no_po_frame(
        filter_template_no_po(result_df),
        no_po_headers,
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_BO_REPORT
    _write_bo_report_workbook(
        guide_path=guide_path,
        output_path=output_path,
        result_df=result_df,
        template_headers=template_headers,
        output_headers=output_headers,
        no_po_df=no_po_df,
        no_po_headers=no_po_headers,
    )

    preview_cols = [
        c
        for c in [
            "Sales document",
            "Material No",
            "Order Quantity",
            "Class",
            "Purchasing Document",
            "Total Price",
            "Milestone",
            "SO Date",
            "Action",
        ]
        if c in result_df.columns
    ]
    preview_rows = (
        result_df[preview_cols]
        .head(15)
        .fillna("")
        .astype(str)
        .replace({"None": "", "nan": "", "<NA>": ""})
        .to_dict(orient="records")
        if preview_cols
        else []
    )

    return BoGuideResult(
        job_id=JOB_ID_BO_GUIDE,
        row_count=len(result_df),
        excluded_qty_equal_count=excluded_qty_equal_count,
        excluded_delivery_charge_count=excluded_delivery_charge_count,
        excluded_rejection_count=excluded_rejection_count,
        excluded_exact_duplicate_count=excluded_exact_duplicate_count,
        excluded_duplicate_count=excluded_duplicate_count,
        source_item_files=source_files,
        guide_path=str(guide_path),
        files={"bo_report": output_path},
        preview=preview_rows,
        missing_sources=missing_sources,
    )
